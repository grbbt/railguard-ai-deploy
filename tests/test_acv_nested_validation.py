"""Synthetic checks for ACV case isolation and the frozen two-feature ranker."""
from datetime import datetime, timedelta
import hashlib
import json

import joblib
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from backend.ps3 import acv


def _workbook(path, hot="01", mode="Full Cooling", cars=8):
    book = Workbook()
    sheet = book.active
    fields = ["Indoor Average Temperature", "ACV Control Temperature (Cooling)",
              "ACV Running Mode", "ACV Setting Mode", "ACV Information Valid"]
    sheet.append(["Time"] + [f"Car {car:02} - {field}" for car in range(1, cars + 1) for field in fields])
    for sample in range(8):
        row = [datetime(2023, 1, 1) + timedelta(seconds=30 * sample)]
        for car in range(1, cars + 1):
            cabin = 28 + sample / 100 if f"{car:02}" == hot else 23 + car / 20 + sample / 100
            row.extend([cabin, 24, mode, "Centralized Control", "Valid"])
        sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def _case(file_id, group):
    cars = [f"{car:02}" for car in range(1, 9)]
    return acv.Case(file_id, 8, pd.DataFrame(.5, index=cars, columns=acv.FEATURES),
                    {car: True for car in cars}, {}, [], [], 30., group)


def test_inner_and_outer_fits_receive_only_training_labels_and_exclude_groups(monkeypatch):
    cases = [_case("a.xlsx", "group-a"), _case("a-copy.xlsx", "group-a"),
             _case("b.xlsx", "group-b"), _case("c.xlsx", "group-c"), _case("d.xlsx", "group-d")]
    groups = {case.file_id: case.content_sha256 for case in cases}
    labels = {case.file_id: "01" for case in cases}
    labels["not-an-input.xlsx"] = "08"
    fits = []

    def fit(name, training, supplied_labels):
        ids = {case.file_id for case in training}
        assert ids == set(supplied_labels), "Held-out or extraneous labels reached model fitting"
        assert supplied_labels == {name: labels[name] for name in ids}
        fits.append(ids)
        return {"name": name, "training": ids}

    def rank(model, heldout):
        assert groups[heldout.file_id] not in {groups[name] for name in model["training"]}
        ranking = list(heldout.features.index)
        return ranking, {car: float(8 - index) for index, car in enumerate(ranking)}

    monkeypatch.setattr(acv, "_fit", fit)
    monkeypatch.setattr(acv, "_ranking", rank)
    _, candidates, nested = acv.evaluate_cases(cases, labels, groups, ["first", "second"])
    assert fits and any(len(ids) < len(cases) - 1 for ids in fits)
    assert all(candidate["ranking_metrics"]["cases"] == len(cases) for candidate in candidates)
    assert nested["ranking_metrics"] == {"cases": 5, "top1_correct": 5, "top2_correct": 5, "mean_rank": 1.0}
    for outer in nested["folds"]:
        assert outer["selected_model"] == "first", "Ties must preserve the declared candidate order"
        outer_group = {case.file_id for case in cases if groups[case.file_id] == groups[outer["file_id"]]}
        assert set(outer["heldout_files"]) == outer_group
        assert not outer_group.intersection(outer["training_files"])
        for inner in outer["inner_folds"]:
            assert not outer_group.intersection(inner["training_files"])
            assert not outer_group.intersection(inner["heldout_files"])
            assert not set(inner["heldout_files"]).intersection(inner["training_files"])


def test_outer_answer_cannot_change_its_inner_selection_or_prediction(monkeypatch):
    cases = [_case(f"{letter}.xlsx", letter) for letter in "abcd"]
    groups = {case.file_id: case.content_sha256 for case in cases}
    labels = {case.file_id: "01" for case in cases}

    def fit(name, training, supplied_labels):
        return {"name": name, "label_sum": sum(int(label) for label in supplied_labels.values())}

    def rank(model, heldout):
        cars = list(heldout.features.index)
        offset = (model["label_sum"] + int(model["name"] == "second")) % 8
        ranked = cars[offset:] + cars[:offset]
        return ranked, {car: float(8 - i) for i, car in enumerate(ranked)}

    monkeypatch.setattr(acv, "_fit", fit)
    monkeypatch.setattr(acv, "_ranking", rank)
    _, _, first = acv.evaluate_cases(cases, labels, groups, ["first", "second"])
    labels["a.xlsx"] = "08"
    _, _, changed = acv.evaluate_cases(cases, labels, groups, ["first", "second"])
    before = next(row for row in first["folds"] if row["file_id"] == "a.xlsx")
    after = next(row for row in changed["folds"] if row["file_id"] == "a.xlsx")
    assert before["inner_candidates"] == after["inner_candidates"]
    assert before["selected_model"] == after["selected_model"]
    assert before["ranked_cars"] == after["ranked_cars"]
    assert before["true_car_rank"] != after["true_car_rank"]


def test_duplicate_parsed_workbooks_with_changed_excel_metadata_stay_together(tmp_path, monkeypatch):
    data, artifacts = tmp_path / "data", tmp_path / "artifacts"
    labels = []
    for car in range(1, 4):
        filename = f"case{car}.xlsx"
        _workbook(data / "Train" / filename, hot=f"{car:02}")
        labels.append({"filename": filename, "faulty_car": f"{car:02}"})
    original, duplicate = data / "Train" / "case1.xlsx", data / "Train" / "changed-package.xlsx"
    book = load_workbook(original)
    book.properties.creator = "A different workbook metadata value"
    book.properties.description = "Telemetry was left unchanged"
    book.save(duplicate)
    book.close()
    assert hashlib.sha256(original.read_bytes()).digest() != hashlib.sha256(duplicate.read_bytes()).digest()
    assert acv.load_case(original).content_sha256 == acv.load_case(duplicate).content_sha256
    labels.append({"filename": duplicate.name, "faulty_car": "01"})
    pd.DataFrame(labels).to_csv(data / "Train_Labels.csv", index=False)
    monkeypatch.setattr(acv, "_candidates", lambda: [acv.BASELINE])
    metadata = acv.train(data, artifacts)
    assert len(set(metadata["input_groups"].values())) == 3
    assert metadata["validation"]["ranking_metrics"]["cases"] == 4
    assert "4 labelled case files" in metadata["limitations"][0]
    for outer in metadata["validation"]["nested_validation"]["folds"]:
        excluded = set(outer["heldout_files"])
        if outer["file_id"] in {original.name, duplicate.name}:
            assert excluded == {original.name, duplicate.name}
        assert not excluded.intersection(outer["training_files"])
        for inner in outer["inner_folds"]:
            assert not excluded.intersection(inner["training_files"] + inner["heldout_files"])
            if original.name in inner["heldout_files"] or duplicate.name in inner["heldout_files"]:
                assert set(inner["heldout_files"]) == {original.name, duplicate.name}
    labels[-1]["faulty_car"] = "02"
    pd.DataFrame(labels).to_csv(data / "Train_Labels.csv", index=False)
    with pytest.raises(ValueError, match="conflicting faulty-car labels"):
        acv.train(data, artifacts)


@pytest.mark.parametrize("mode", ["Not Cooling", "Cooling Disabled", "Unknown cooling", "Full Cooling disabled"])
def test_negated_or_unrecognised_cooling_text_is_excluded(tmp_path, mode):
    path = tmp_path / "unknown-mode.xlsx"
    _workbook(path, mode=mode)
    case = acv.load_case(path)
    assert not any(case.available.values())
    assert case.features.isna().all().all()
    assert all(detail["usable_cooling_rows"] == 0 for detail in case.details.values())
    assert any("unrecognised cooling mode text" in warning for warning in case.warnings)


@pytest.mark.parametrize("mode", [" FULL COOLING ", "Half Cooling", "Automatic Cooling"])
def test_documented_cooling_modes_remain_usable(tmp_path, mode):
    path = tmp_path / "known-mode.xlsx"
    _workbook(path, mode=mode)
    case = acv.load_case(path)
    assert all(case.available.values())
    assert all(detail["usable_cooling_rows"] == 8 for detail in case.details.values())


@pytest.mark.parametrize("cars", [7, 9])
def test_non_eight_car_workbook_fails_before_inference(tmp_path, cars):
    path = tmp_path / "wrong-formation.xlsx"
    _workbook(path, cars=cars)
    with pytest.raises(ValueError, match="exactly eight distinct cars"):
        acv.load_case(path)


def test_two_feature_candidate_saves_a_compatible_frozen_predictor(tmp_path, monkeypatch):
    data, artifacts = tmp_path / "data", tmp_path / "artifacts"
    labels = []
    for car in range(1, 5):
        name = f"case{car}.xlsx"
        _workbook(data / "Train" / name, hot=f"{car:02}")
        labels.append({"filename": name, "faulty_car": f"{car:02}"})
    pd.DataFrame(labels).to_csv(data / "Train_Labels.csv", index=False)
    (data / "Test").mkdir()
    (data / "Test" / "never-read.xlsx").write_bytes(b"Not a workbook: training must never read Test")
    real_loader = acv.load_case

    def permitted_loader(path):
        assert path.parent.name != "Test", "Training attempted to read hidden Test inputs"
        return real_loader(path)

    monkeypatch.setattr(acv, "load_case", permitted_loader)
    monkeypatch.setattr(acv, "_candidates", lambda: [acv.COOLING])
    metadata = acv.train(data, artifacts)
    assert metadata["model_name"] == acv.COOLING
    assert metadata["selected_model_features"] == ["above_peer_fraction", "above_target_fraction"]
    bundle = joblib.load(artifacts / "model.joblib")
    assert list(bundle["preprocess"].feature_names_in_) == metadata["selected_model_features"]
    assert bundle["estimator"].coef_.shape == (1, 2)
    assert metadata["validation"]["ranking_metrics"]["cases"] == 4
    assert metadata["validation"]["nested_validation"]["ranking_metrics"]["cases"] == 4
    assert metadata["validation"]["ranking_metrics"]["top1_correct"] == 4
    model_bytes_before = (artifacts / "model.joblib").read_bytes()

    def forbidden_fit(*args, **kwargs):
        pytest.fail("Inference attempted to refit the model")

    monkeypatch.setattr(acv, "_fit", forbidden_fit)
    probe = tmp_path / "upload.xlsx"
    _workbook(probe, hot="08")
    report = acv.predict(probe, artifacts)
    assert report["prediction_rows"][0]["ranked_cars"].split("|")[0] == "08"
    assert (artifacts / "model.joblib").read_bytes() == model_bytes_before
    evidence = {item["id"]: item for item in report["evidence"]}
    assert "Observed model input" in evidence["acv-above_peer_fraction-08"]["detail"]
    assert "not a selected-model input" in evidence["acv-peer_q90_rank-08"]["detail"]
    json.dumps(report, allow_nan=False)


def test_nested_evaluation_refuses_fewer_than_three_distinct_groups():
    cases = [_case("first.xlsx", "same"), _case("copy.xlsx", "same"), _case("second.xlsx", "other")]
    with pytest.raises(ValueError, match="at least three distinct input groups"):
        acv.evaluate_cases(cases, {case.file_id: "01" for case in cases},
                           {case.file_id: case.content_sha256 for case in cases}, [acv.BASELINE])
