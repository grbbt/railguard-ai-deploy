"""Targeted regression tests for organiser schemas, grouping and frozen inference."""
import json
import hashlib
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import joblib
from openpyxl import Workbook, load_workbook

from backend.ps3 import acv, door


def _door_fixture(folder: Path, actions=15):
    folder.mkdir(parents=True, exist_ok=True)
    rows, labels = [], []
    for action in range(actions):
        start = pd.Timestamp("2023-07-05") + pd.Timedelta(seconds=action * 10)
        abnormal = action % 3 == 0
        for i in range(12):
            rows.append({"Datetime": door.native_timestamp(start + pd.Timedelta(milliseconds=i * 20)), "Motor current(mA)": 200 + action + i * (15 if abnormal else 2), "Motor Voltage(10mV)": 1200 + i,
                         "Motor electrodynamic force": i * 20 + action, "Door leaf position": 700 - i * 30, "Close command": 1, "Open command": 0, "Door is opening": 0, "Door is closing": 1})
        labels.append({"start_time": rows[-12]["Datetime"], "end_time": rows[-1]["Datetime"], "status": door.LABELS[int(abnormal)], "n_rows": 12})
    pd.DataFrame(rows).to_csv(folder / "Train.csv", index=False)
    pd.DataFrame(labels).to_csv(folder / "Train_Segments_Answer.csv", index=False)
    (folder / "Test.csv").write_text("DO NOT LOAD THIS WHILE TRAINING", encoding="utf-8")


def _acv_fixture(path: Path, hot="01", unavailable=(), rich=False, cadence=30):
    workbook = Workbook()
    sheet = workbook.active
    if rich:
        parameters = ["Passenger Cabin Temperature Detected Value", "Target Temperature Value", "ACV Running Mode", "ACV Control Mode", "Observation Area Temperature Detected Value", "Target Temperature +1K"]
    else:
        parameters = ["Indoor Average Temperature", "ACV Control Temperature (Cooling)", "ACV Running Mode", "ACV Setting Mode", "ACV Information Valid"]
    headers = ["Car model", "Train number", "Time"] + [f"Car {i:02} - {parameter}" for i in range(1, 9) for parameter in parameters]
    sheet.append(headers)
    for n in range(8):
        row = ["fixture", "anonymous", (pd.Timestamp("2023-01-01") + pd.Timedelta(seconds=n * cadence)).to_pydatetime()]
        for i in range(1, 9):
            car = f"{i:02}"
            if car in unavailable:
                row.extend(["None"] * len(parameters))
            else:
                row.extend([24 + (4 if car == hot else i / 20) + n / 100, 24, "Full Cooling", "Centralized Control"])
                row.extend([-50, 1] if rich else ["Valid"])
        sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def test_native_door_milliseconds_are_integer_fields():
    assert door.parse_timestamp("2023-7-5-0-0-0-20").microsecond == 20_000
    assert door.parse_timestamp("2023-7-5-0-0-0-1").microsecond == 1_000
    assert door.native_timestamp(door.parse_timestamp("2023-7-5-0-0-0-760")) == "2023-7-5-0-0-0-760"
    with pytest.raises(ValueError, match="millisecond"):
        door.parse_timestamp("2023-7-5-0-0-0-1000")


def test_door_iou_scores_wrong_labels_duplicates_and_greedy_matching():
    def segment(a, b, label="Normal"):
        return {"start_time": f"2023-01-01T00:00:{a:02}", "end_time": f"2023-01-01T00:00:{b:02}", "prediction": label}
    true = [segment(0, 10)]
    assert door.iou_weighted_f1(true, [segment(0, 10)]) == 1
    assert door.iou_weighted_f1(true, [segment(0, 5)]) == .5
    assert door.iou_weighted_f1(true, [segment(0, 10, "Abnormal resistance")]) == 0
    assert door.iou_weighted_f1(true, [segment(0, 10), segment(0, 10)]) == pytest.approx(2 / 3)
    assert door.iou_weighted_f1([segment(0, 10), segment(6, 16)], [segment(0, 16), segment(0, 5)]) == .3125


def test_door_blocks_keep_duplicate_actions_out_of_training():
    answers = pd.DataFrame(index=range(10))
    groups = [str(i) for i in range(10)]
    groups[8] = groups[1]
    folds = door.blocked_folds(answers, groups=groups)
    for training, heldout in folds:
        assert np.all(np.diff(heldout) == 1)
        assert not ({groups[i] for i in training} & {groups[i] for i in heldout})
    assert sorted(np.concatenate([v for _, v in folds])) == list(range(10))


def test_door_schema_duplicate_clock_and_no_label_feature(tmp_path):
    _door_fixture(tmp_path)
    frame = door.load_stream(tmp_path / "Train.csv")
    assert len(door.segment_stream(frame, door.segmentation_config(frame))) == 15
    assert "status" not in door.action_features(frame.iloc[:12])
    raw = pd.read_csv(tmp_path / "Train.csv")
    raw.loc[1, "Datetime"] = raw.loc[0, "Datetime"]
    raw.to_csv(tmp_path / "duplicate.csv", index=False)
    with pytest.raises(ValueError, match="unique and increasing"):
        door.load_stream(tmp_path / "duplicate.csv")


def test_door_training_ignores_test_and_prediction_never_refits(tmp_path, monkeypatch):
    data, artifacts = tmp_path / "data", tmp_path / "artifacts"
    _door_fixture(data)
    monkeypatch.setattr(door, "_candidate_names", lambda: ["Regularized logistic regression", "Normal-cycle Isolation Forest"])
    metadata = door.train(data, artifacts)
    assert metadata["training_rows"] == 180
    assert metadata["model_sha256"] == hashlib.sha256((artifacts / "model.joblib").read_bytes()).hexdigest()
    assert json.loads((artifacts / "metadata.json").read_text(encoding="utf-8"))["model_sha256"] == metadata["model_sha256"]
    assert not list(artifacts.glob(".*.tmp"))
    assert all(not (set(f["training_actions"]) & set(f["heldout_actions"])) for f in metadata["validation"]["folds"])
    def forbidden(*args, **kwargs):
        raise AssertionError("Inference attempted to fit")
    monkeypatch.setattr(door, "_fit", forbidden)
    report = door.predict(data / "Train.csv", artifacts)
    assert len(report["prediction_rows"]) == 15
    assert set(report["prediction_rows"][0]) == {"start_time", "end_time", "prediction"}
    assert max(len(series["points"]) for series in report["series"]) <= 400
    json.dumps(report, allow_nan=False)


def test_acv_modes_missing_cars_and_suspected_sentinel_are_not_faults(tmp_path):
    path = tmp_path / "rich.xlsx"
    _acv_fixture(path, unavailable=("05", "06", "07", "08"), rich=True, cadence=10)
    case = acv.load_case(path)
    assert case.cadence == 10
    assert sum(case.available.values()) == 4
    assert case.details["01"]["residual_median"] > 3
    assert any("suspected sentinel" in warning for warning in case.warnings)
    assert "Target Temperature +1K" not in case.features.columns
    order, scores = acv._ranking({"name": acv.BASELINE}, case)
    assert order[0] == "01"
    assert order[-4:] == ["05", "06", "07", "08"]
    assert all(scores[car] is None for car in order[-4:])
    assert set(order) == {f"{i:02}" for i in range(1, 9)}


def test_acv_all_missing_evidence_is_explicit_not_healthy(tmp_path):
    path = tmp_path / "missing.xlsx"
    _acv_fixture(path, unavailable=tuple(f"{i:02}" for i in range(1, 9)))
    case = acv.load_case(path)
    assert not any(case.available.values())
    assert any("entirely identifier-based" in warning for warning in case.warnings)
    ranks, scores = acv._ranking({"name": acv.BASELINE}, case)
    assert all(score is None for score in scores.values())
    assert ranks == [f"{i:02}" for i in range(1, 9)]


def test_acv_report_exposes_discriminating_inputs_when_medians_tie(tmp_path):
    path = tmp_path / "tied-medians.xlsx"
    _acv_fixture(path, unavailable=("08",))
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    for car in range(1, 8):
        column = headers.index(f"Car {car:02} - Indoor Average Temperature") + 1
        for row in range(2, 10):
            sheet.cell(row, column).value = 24 + (2 if row >= (7 if car == 1 else 9) else 0)
    workbook.save(path)
    workbook.close()
    artifacts = tmp_path / "artifacts"
    artifacts.mkdir()
    joblib.dump({"name": acv.BASELINE, "version": acv.VERSION, "subsystem": "acv", "feature_names": acv.FEATURES}, artifacts / "model.joblib")
    (artifacts / "metadata.json").write_text(json.dumps({
        "subsystem": "acv", "artifact_version": acv.VERSION, "model_name": acv.BASELINE,
        "feature_names": acv.FEATURES, "model_sha256": hashlib.sha256((artifacts / "model.joblib").read_bytes()).hexdigest(),
    }), encoding="utf-8")
    case = acv.load_case(path)
    report = acv.predict(path, artifacts)
    evidence = {entry["id"]: entry for entry in report["evidence"]}
    assert len(evidence) <= 40
    assert evidence["acv-car-01"]["value"] == evidence["acv-car-02"]["value"] == 0
    assert evidence["acv-above_target_fraction-01"]["value"] == 37.5
    assert evidence["acv-above_target_fraction-02"]["value"] == 12.5
    for feature in ("above_target_fraction", "above_peer_fraction", "peer_q90_rank"):
        assert evidence[f"acv-{feature}-01"]["value"] == pytest.approx(case.features.loc["01", feature] * 100)
        assert evidence[f"acv-{feature}-08"]["value"] is None
        assert "not a probability or feature-contribution estimate" in evidence[f"acv-{feature}-01"]["detail"]
    assert report["entities"][-1]["status"] == "unavailable"
    json.dumps(report, allow_nan=False)


def test_acv_invalid_clock_and_duplicate_headers_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    workbook = Workbook()
    workbook.active.append(["Time", "Time"])
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="unique"):
        acv.load_case(path)
    _acv_fixture(path)
    workbook = load_workbook(path)
    workbook.active.cell(3, 3).value = workbook.active.cell(2, 3).value
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="timestamps must be valid, unique and increasing"):
        acv.load_case(path)


def test_acv_corrupt_workbook_is_a_useful_file_error(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"Not an Excel archive")
    with pytest.raises(ValueError, match="Cannot read ACV Excel workbook"):
        acv.load_case(path)


def test_acv_invalid_states_and_formula_values_cannot_be_sensor_evidence(tmp_path):
    path = tmp_path / "states.xlsx"
    _acv_fixture(path)
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, headers.index("Car 01 - ACV Information Valid") + 1).value = "Invalid"
        sheet.cell(row, headers.index("Car 02 - Indoor Average Temperature") + 1).value = "=999"
    workbook.save(path)
    workbook.close()
    case = acv.load_case(path)
    assert not case.available["01"] and not case.available["02"]
    assert case.features.loc[["01", "02"]].isna().all().all()
    assert any("formula cells" in warning for warning in case.warnings)


def test_acv_duplicate_workbooks_stay_in_the_same_fold_and_keep_file_ids(tmp_path, monkeypatch):
    data, artifacts = tmp_path / "data", tmp_path / "artifacts"
    labels = []
    for n in range(3):
        name = f"case{n}.xlsx"
        _acv_fixture(data / "Train" / name, hot=f"{n + 1:02}")
        labels.append({"filename": name, "faulty_car": f"{n + 1:02}"})
    shutil.copyfile(data / "Train" / "case0.xlsx", data / "Train" / "duplicate.xlsx")
    labels.append({"filename": "duplicate.xlsx", "faulty_car": "01"})
    pd.DataFrame(labels).to_csv(data / "Train_Labels.csv", index=False)
    monkeypatch.setattr(acv, "_candidates", lambda: [acv.BASELINE])
    metadata = acv.train(data, artifacts)
    folds = metadata["validation"]["folds"]
    assert {fold["file_id"] for fold in folds} == {row["filename"] for row in labels}
    for fold in folds:
        assert not (set(fold["training_files"]) & set(fold["heldout_files"]))
        if fold["file_id"] in ("case0.xlsx", "duplicate.xlsx"):
            assert set(fold["heldout_files"]) == {"case0.xlsx", "duplicate.xlsx"}
    labels[-1]["faulty_car"] = "02"
    pd.DataFrame(labels).to_csv(data / "Train_Labels.csv", index=False)
    with pytest.raises(ValueError, match="conflicting faulty-car labels"):
        acv.train(data, artifacts)


def test_acv_exact_rank_decay_not_reciprocal_rank():
    ranking = [f"{i:02}" for i in range(1, 9)]
    assert acv.rank_decay(ranking, "02", 8) == .875
    assert acv.rank_decay(ranking, "08", 8) == .125
    assert acv.rank_decay(ranking, "09", 8) == 0


def test_acv_leave_case_out_artifacts_and_frozen_predictions(tmp_path, monkeypatch):
    data, artifacts = tmp_path / "data", tmp_path / "artifacts"
    labels = []
    for n in range(3):
        name = f"case{n}.xlsx"
        _acv_fixture(data / "Train" / name, hot=f"{n + 1:02}")
        labels.append({"filename": name, "faulty_car": f"{n + 1:02}"})
    pd.DataFrame(labels).to_csv(data / "Train_Labels.csv", index=False)
    (data / "Test").mkdir()
    (data / "Test" / "poison.xlsx").write_bytes(b"NEVER READ TEST DURING TRAINING")
    monkeypatch.setattr(acv, "_candidates", lambda: [acv.BASELINE, acv.PAIRWISE, acv.FOREST])
    metadata = acv.train(data, artifacts)
    assert metadata["training_files"] == 3
    assert metadata["model_sha256"] == hashlib.sha256((artifacts / "model.joblib").read_bytes()).hexdigest()
    assert json.loads((artifacts / "metadata.json").read_text(encoding="utf-8"))["model_sha256"] == metadata["model_sha256"]
    assert not list(artifacts.glob(".*.tmp"))
    assert metadata["validation"]["score"] == 1
    for fold in metadata["validation"]["folds"]:
        assert fold["file_id"] not in fold["training_files"]
    def forbidden(*args, **kwargs):
        raise AssertionError("Inference attempted to fit")
    monkeypatch.setattr(acv, "_fit", forbidden)
    report = acv.predict(data / "Train" / "case1.xlsx", artifacts)
    assert report["prediction_rows"][0]["ranked_cars"].startswith("02|")
    assert set(report["prediction_rows"][0]) == {"file_id", "ranked_cars"}
    assert all(entity["status"] == "ranked" for entity in report["entities"])
    assert max(len(series["points"]) for series in report["series"]) <= 400
    json.dumps(report, allow_nan=False)
