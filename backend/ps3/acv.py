"""Case-level refrigerant-leak ranking with explicit missing-evidence handling.

Every learned row is one car's whole-case summary. Validation excludes an entire
workbook, and the case label never labels a timestamp or a fault-onset interval.
"""
from __future__ import annotations

import hashlib
import json
import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sklearn.ensemble import IsolationForest
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

from backend.ps3.rail import _load_model
from backend.ps3.preview import sampled_trace_points

VERSION = 2
BASELINE = "Cooling peer-deviation ranking"
PAIRWISE = "Pairwise regularized logistic ranking"
COOLING = "Pairwise cooling-exceedance ranking"
COOLING_FEATURES = ["above_peer_fraction", "above_target_fraction"]
COOLING_MODES = {"automatic cooling", "full cooling", "half cooling"}
FOREST = "Normal-car Isolation Forest"
BOOST = "XGBoost pairwise ranking"
ALIASES = {
    "cabin": ("Indoor Average Temperature", "Passenger Cabin Temperature Detected Value"),
    "target": ("ACV Control Temperature (Cooling)", "Target Temperature Value"),
    "outside": ("Outdoor Average Temperature", "Outside Temperature Sensor Reading", "Fresh Air Temperature Detected Value"),
    "running": ("ACV Running Mode",),
    "valid": ("ACV Information Valid",),
    "control": ("ACV Setting Mode", "ACV Control Mode"),
    "operating": ("ACV Operating Mode",),
    "observation": ("Observation Area Temperature Detected Value",),
}
FEATURES = ["peer_median_rank", "peer_q90_rank", "residual_median_rank", "residual_q90_rank", "residual_spread_rank", "above_peer_fraction", "above_target_fraction", "warming_fraction"]
LIMITATIONS = [
    "The released labelled case set is small; case independence and acquisition grouping are unverified.",
    "Leave-one-case-out scores were used for candidate selection; they are not independent official test scores.",
    "One faulty car per case is the organiser task assumption. A ranking does not verify a refrigerant leak or locate its onset.",
    "Temperature and pressure scaling are undocumented; evidence is reported in raw source units, not calibrated physical units.",
    "Cars without usable cooling/cabin/target evidence are placed last in stable identifier order solely to satisfy the required complete ranking; their condition is unknown.",
]


@dataclass
class Case:
    file_id: str
    rows: int
    features: pd.DataFrame
    available: dict[str, bool]
    details: dict[str, dict]
    series: list[dict]
    warnings: list[str]
    cadence: float
    content_sha256: str = ""


def _sha(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def _numeric(series: pd.Series) -> pd.Series:
    values = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
    if values.abs().gt(1e12).any():
        raise ValueError("ACV sensor values exceed the supported numeric range; check raw source units and corrupted cells.")
    return values


def _weighted(values: pd.Series, weights: np.ndarray, q: float = .5) -> float:
    a = values.to_numpy(float)
    good = np.isfinite(a) & np.isfinite(weights) & (weights > 0)
    if not good.any():
        return np.nan
    order = np.argsort(a[good], kind="stable")
    v, w = a[good][order], weights[good][order]
    return float(v[min(np.searchsorted(np.cumsum(w), q * w.sum()), len(v) - 1)])


def _fraction(values: pd.Series, weights: np.ndarray) -> float:
    a = values.to_numpy(float)
    good = np.isfinite(a)
    return float(np.average(a[good] > 0, weights=weights[good])) if good.any() else np.nan


def _finite(value):
    return float(value) if value is not None and np.isfinite(value) else None


def load_case(path: Path) -> Case:
    """Stream the workbook, retaining only explicitly recognised sensor/state columns."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("ACV requires an .xlsx case workbook with Time and Car NN - parameter headers.")
    try:
        with ZipFile(path) as archive:
            expanded_bytes = sum(item.file_size for item in archive.infolist())
    except (OSError, BadZipFile) as exc:
        raise ValueError("Cannot read ACV Excel workbook; supply a readable, uncorrupted .xlsx file.") from exc
    if expanded_bytes > 512 * 1024 * 1024:
        raise ValueError("ACV workbook expands beyond the 512 MiB inspection limit.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (OSError, BadZipFile, KeyError, InvalidFileException, SyntaxError, ValueError, TypeError, OverflowError) as exc:
        raise ValueError("Cannot read ACV Excel workbook; its workbook structure or cell values are malformed.") from exc
    warnings = []
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("ACV expects one case on one worksheet; combine or select the intended sheet before upload.")
        sheet = workbook.worksheets[0]
        if sheet.max_column > 800 or sheet.max_row > 250_001 or sheet.max_column * sheet.max_row > 20_000_000:
            raise ValueError("ACV worksheet exceeds the supported 800 columns / 250,000 rows / 20 million cells.")
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            raise ValueError("ACV workbook is empty.")
        headers = [str(value).strip() if value is not None else "" for value in headers]
        if len({header.casefold() for header in headers}) != len(headers) or any(not value for value in headers):
            raise ValueError("ACV headers must be nonblank and unique, ignoring case and surrounding spaces.")
        time_columns = [i for i, value in enumerate(headers) if value.lower() in ("time", "datetime", "timestamp")]
        if len(time_columns) > 1:
            raise ValueError("ACV workbook has ambiguous time columns; retain exactly one Time, Datetime or Timestamp column.")
        time_column = next(iter(time_columns), None)
        if time_column is None:
            raise ValueError("ACV workbook is missing Time/Datetime.")
        columns = {}
        all_cars = set()
        for i, header in enumerate(headers):
            match = re.fullmatch(r"Car (\d{2}) - (.+)", header, flags=re.IGNORECASE)
            if match:
                car, parameter = match.groups()
                all_cars.add(car)
                columns[(car, parameter.casefold())] = (i, header)
        cars = sorted(all_cars)
        if len(cars) != 8:
            raise ValueError("ACV requires exactly eight distinct cars for the PS3 ranking format; use each workbook's Car NN headers.")
        chosen = {}
        selected_indexes = {time_column}
        for car in cars:
            chosen[car] = {}
            for role, aliases in ALIASES.items():
                if sum((car, alias.casefold()) in columns for alias in aliases) > 1:
                    raise ValueError(f"ACV Car {car} has multiple aliases for {role}; retain one unambiguous sensor column.")
                column = next((columns[(car, alias.casefold())] for alias in aliases if (car, alias.casefold()) in columns), None)
                if column:
                    chosen[car][role] = column
                    selected_indexes.add(column[0])
        selected_indexes = sorted(selected_indexes)
        selected_rows = []
        formulas = 0
        for row in iterator:
            if all(value is None for value in row):
                continue
            values = []
            for index in selected_indexes:
                value = row[index] if index < len(row) else None
                if isinstance(value, str) and value.startswith("="):
                    value = None; formulas += 1
                values.append(value)
            selected_rows.append(values)
            if len(selected_rows) > 250_000:
                raise ValueError("ACV workbook has more than 250,000 data rows.")
    except (ValueError, TypeError, IndexError, KeyError, OverflowError, SyntaxError, BadZipFile, OSError) as exc:
        if isinstance(exc, ValueError) and str(exc).startswith("ACV "):
            raise
        raise ValueError("Malformed ACV workbook; check worksheet dimensions and cell contents.") from exc
    finally:
        workbook.close()
    if len(selected_rows) < 3:
        raise ValueError("ACV requires at least three timestamped readings.")
    raw = pd.DataFrame(selected_rows, columns=[headers[i] for i in selected_indexes])
    if formulas:
        warnings.append(f"{formulas} formula cells in selected channels were treated as unavailable; formulas are never executed.")
    times_raw = raw[headers[time_column]]
    if times_raw.map(lambda v: isinstance(v, (int, float))).any():
        raise ValueError("ACV Time must contain Excel datetimes or date strings, not unformatted numeric date serials.")
    try:
        times = pd.to_datetime(times_raw, errors="coerce", format="mixed")
        if times.isna().any() or times.duplicated().any() or not times.is_monotonic_increasing:
            raise ValueError("ACV timestamps must be valid, unique and increasing.")
        dt = times.diff().dt.total_seconds()
    except (TypeError, AttributeError) as exc:
        raise ValueError("ACV timestamps must use one consistent timezone-free clock.") from exc
    if times.dt.tz is not None:
        raise ValueError("ACV expects the timezone-free source clock; do not add a guessed timezone.")
    cadence = float(dt.dropna().mode().iloc[0])
    if not np.isfinite(cadence) or cadence <= 0:
        raise ValueError("Cannot determine ACV recorded cadence.")
    # Do not assign the last observed value to long periods with no observation.
    next_dt = times.shift(-1).sub(times).dt.total_seconds().to_numpy(float)
    weights = np.where(np.isfinite(next_dt) & (next_dt > 0) & (next_dt <= 3 * cadence), next_dt, cadence)
    if (dt > 3 * cadence).any():
        warnings.append("Acquisition gaps are present; cooling summaries do not interpolate or carry temperatures across long gaps.")
    if cadence != 30:
        warnings.append(f"Observed modal cadence is {cadence:g} s; features use this file's timing rather than assuming 30 s.")
    residual, masks, modes, signals, targets, details = {}, {}, {}, {}, {}, {}
    preview_peer_gaps = {}
    empty = pd.Series(np.nan, index=raw.index)
    for car in cars:
        lookup = chosen[car]
        series = {role: raw[column[1]] for role, column in lookup.items()}
        cabin = _numeric(series.get("cabin", empty))
        target = _numeric(series.get("target", empty))
        run = series.get("running", empty).astype(str).str.strip().str.casefold()
        cooling = run.isin(COOLING_MODES)
        unknown_cooling = run.str.contains("cooling", regex=False, na=False) & ~cooling
        if unknown_cooling.any():
            warnings.append(f"Car {car}: {int(unknown_cooling.sum())} readings use unrecognised cooling mode text and are excluded from cooling evidence.")
        valid = pd.Series(True, index=raw.index)
        if "valid" in series:
            valid &= series["valid"].astype(str).str.strip().str.casefold().eq("valid")
        if "control" in series:
            valid &= ~series["control"].astype(str).str.strip().str.casefold().isin(["invalid", "none", "nan", ""])
        mask = valid & cooling & cabin.notna() & target.notna()
        masks[car], modes[car] = mask, run
        residual[car] = (cabin - target).where(mask)
        signals[car] = cabin.where(valid)
        targets[car] = target
        details[car] = {"usable_cooling_rows": int(mask.sum()), "total_rows": len(raw), "cabin_column": lookup.get("cabin", (None, None))[1], "target_column": lookup.get("target", (None, None))[1], "running_column": lookup.get("running", (None, None))[1], "numeric_cabin_rows": int(cabin.notna().sum()), "numeric_target_rows": int(target.notna().sum())}
        if "observation" in series:
            sentinel_count = int(_numeric(series["observation"]).eq(-50).sum())
            if sentinel_count:
                warnings.append(f"Car {car}: {sentinel_count} observation-area values equal -50, a suspected sentinel. That channel is excluded; cabin measurements are used without interpreting -50 as a fault.")
        if "operating" in series and series["operating"].astype(str).str.casefold().eq("invalid").any():
            warnings.append(f"Car {car}: Operating Mode includes Invalid while other channels may remain usable. Features use explicit Information Valid where present, valid control states and named Cooling running modes; operating-state meaning remains uncertain.")
    residual_frame = pd.DataFrame(residual)
    raw_features = []
    available = {}
    for car in cars:
        other = [value for value in cars if value != car]
        peer_values = residual_frame[other]
        matched = pd.DataFrame({peer: residual_frame[peer].where(modes[peer] == modes[car]) for peer in other})
        same_mode_reference = matched.median(axis=1).where(matched.notna().sum(axis=1) >= 2)
        fallback = peer_values.median(axis=1).where(peer_values.notna().sum(axis=1) >= 2)
        reference = same_mode_reference.fillna(fallback)
        gap = (residual_frame[car] - reference).where(masks[car])
        preview_peer_gaps[car] = gap
        available[car] = details[car]["usable_cooling_rows"] >= 3
        details[car].update({"peer_comparison_rows": int(gap.notna().sum()), "same_mode_peer_rows": int((same_mode_reference.notna() & masks[car]).sum()), "residual_median": _finite(_weighted(residual_frame[car], weights)), "peer_gap_median": _finite(_weighted(gap, weights)), "peer_gap_q90": _finite(_weighted(gap, weights, .9))})
        # Derivatives are limited to adjacent observations in an unchanged regime.
        change = signals[car].diff().div(dt / 60)
        stable = masks[car] & masks[car].shift(fill_value=False) & modes[car].eq(modes[car].shift()) & targets[car].eq(targets[car].shift()) & dt.le(3 * cadence)
        change = change.where(stable)
        own = residual_frame[car]
        raw_features.append({"car": car, "peer_median": _weighted(gap, weights), "peer_q90": _weighted(gap, weights, .9), "residual_median": _weighted(own, weights), "residual_q90": _weighted(own, weights, .9), "residual_spread": _weighted(own, weights, .9) - _weighted(own, weights, .1), "above_peer_fraction": _fraction(gap, weights), "above_target_fraction": _fraction(own, weights), "warming_fraction": _fraction(change, weights)})
    summaries = pd.DataFrame(raw_features).set_index("car")
    features = pd.DataFrame(index=cars, columns=FEATURES, dtype=float)
    for name in ("peer_median", "peer_q90", "residual_median", "residual_q90", "residual_spread"):
        features[f"{name}_rank"] = summaries[name].where(pd.Series(available)).rank(method="average", pct=True)
    for name in ("above_peer_fraction", "above_target_fraction", "warming_fraction"):
        features[name] = summaries[name]
    for car in cars:
        if not available[car]:
            features.loc[car] = np.nan
            warnings.append(f"Car {car}: insufficient valid cooling/cabin/target observations; condition is unknown and any output tail position is administrative, not a healthy diagnosis.")
    if not any(available.values()):
        warnings.append("No car has sufficient usable cooling evidence. The required ordering is entirely identifier-based and cannot localise a leak.")
    elapsed = times.sub(times.iloc[0]).dt.total_seconds() / 3600
    plots = []
    for car in cars:
        for label, values in [
            ("cabin measurement", signals[car]),
            ("cooling target", targets[car].where(masks[car])),
            ("cabin-minus-target residual", residual_frame[car]),
            ("residual above cooling peers", preview_peer_gaps[car]),
        ]:
            if values.notna().any():
                plots.append({"name": f"Car {car} {label}", "x_label": "Hours from recording start",
                              "y_label": "Raw temperature units; scaling undocumented",
                              "points": sampled_trace_points(elapsed, values, break_before=dt.gt(3 * cadence))})
    # Group workbooks with identical selected inputs even when Excel packaging,
    # irrelevant sheets metadata, or source filenames differ.
    canonical = raw.rename(columns=lambda name: name.casefold()).sort_index(axis=1).copy()
    canonical[headers[time_column].casefold()] = times.astype(str)
    content_sha256 = hashlib.sha256(canonical.to_json(orient="split", date_format="iso").encode()).hexdigest()
    return Case(path.name, len(raw), features, available, details, plots, list(dict.fromkeys(warnings)), cadence, content_sha256)


def rank_decay(ranked_cars: list[str], faulty_car: str, n: int) -> float:
    if faulty_car not in ranked_cars or n <= 0:
        return 0.
    return max(0., (n - ranked_cars.index(faulty_car)) / n)


def _candidates() -> list[str]:
    result = [BASELINE, PAIRWISE, FOREST, COOLING]
    try:
        import xgboost  # noqa: F401
        result.append(BOOST)
    except ImportError:
        pass
    return result


def _fit(name: str, cases: list[Case], labels: dict[str, str]) -> dict:
    model = {"name": name, "feature_names": FEATURES, "version": VERSION, "subsystem": "acv"}
    if name == BASELINE:
        return model
    frames, targets, groups = [], [], []
    for case in cases:
        usable = case.features.loc[[car for car in case.features.index if case.available[car]]]
        if labels[case.file_id] not in usable.index or len(usable) < 2:
            continue
        frames.append(usable)
        targets.append(np.asarray([int(car == labels[case.file_id]) for car in usable.index]))
        groups.append(len(usable))
    if not frames:
        raise ValueError("ACV training cases lack usable faulty-car and peer measurements.")
    x, y = pd.concat(frames), np.concatenate(targets)
    if name == COOLING:
        x = x[COOLING_FEATURES]
    preprocess = Pipeline([("impute", SimpleImputer(strategy="median", keep_empty_features=True)), ("scale", StandardScaler())])
    if name == FOREST:
        values = preprocess.fit_transform(x.iloc[np.flatnonzero(y == 0)])
        estimator = IsolationForest(n_estimators=160, contamination="auto", random_state=42, n_jobs=1).fit(values)
    else:
        values = preprocess.fit_transform(x)
        if name in (PAIRWISE, COOLING):
            differences, responses, offset = [], [], 0
            for group in groups:
                group_y = y[offset:offset + group]
                positive = values[offset + int(np.flatnonzero(group_y == 1)[0])]
                for i in np.flatnonzero(group_y == 0):
                    diff = positive - values[offset + i]
                    differences.extend([diff, -diff]); responses.extend([1, 0])
                offset += group
            estimator = LogisticRegression(C=.3, fit_intercept=False, max_iter=2000, random_state=42).fit(np.asarray(differences), responses)
        else:
            from xgboost import XGBRanker
            estimator = XGBRanker(objective="rank:pairwise", n_estimators=80, max_depth=2, learning_rate=.04, min_child_weight=2, reg_lambda=5, subsample=.9, colsample_bytree=.9, random_state=42, n_jobs=1)
            estimator.fit(values, y, group=groups)
    model.update({"preprocess": preprocess, "estimator": estimator})
    return model


def _ranking(model: dict, case: Case) -> tuple[list[str], dict[str, float | None]]:
    usable = [car for car in case.features.index if case.available[car]]
    scores = {car: None for car in case.features.index}
    if usable:
        x = case.features.loc[usable, COOLING_FEATURES if model["name"] == COOLING else FEATURES]
        if model["name"] == BASELINE:
            # Relative, not probabilistic: a persistent warm residual compared
            # with cooling peers outranks a lower residual on the same recording.
            values = (.60 * x.peer_median_rank.fillna(x.residual_median_rank) + .25 * x.peer_q90_rank.fillna(x.residual_q90_rank) + .15 * x.residual_median_rank).to_numpy(float)
        else:
            transformed = model["preprocess"].transform(x)
            if model["name"] == FOREST:
                values = -model["estimator"].score_samples(transformed)
            elif model["name"] in (PAIRWISE, COOLING):
                values = model["estimator"].decision_function(transformed)
            else:
                values = model["estimator"].predict(transformed)
        values = np.asarray(values, dtype=float)
        if values.shape != (len(usable),) or not np.isfinite(values).all():
            raise ValueError("The frozen ACV model returned invalid ranking scores for cars with usable measurements.")
        scores.update({car: float(value) for car, value in zip(usable, values)})
    ranked = sorted(case.features.index, key=lambda car: (scores[car] is None, -scores[car] if scores[car] is not None else 0, car))
    return ranked, scores


def ranking_metrics(rows: list[dict]) -> dict:
    """Separate exact first choice from the organiser's partial-credit metric."""
    if not rows:
        raise ValueError("ACV validation requires held-out case predictions.")
    ranks = [row["true_car_rank"] for row in rows]
    return {"cases": len(rows), "top1_correct": sum(rank == 1 for rank in ranks),
            "top2_correct": sum(rank <= 2 for rank in ranks), "mean_rank": float(np.mean(ranks))}


def evaluate_cases(cases: list[Case], labels: dict[str, str], groups: dict[str, str], names: list[str]) -> tuple[dict, list[dict], dict]:
    """Retrospective nested case validation; never inspect Test or fit held-out cases.

    Outer folds evaluate the selection procedure. Inner folds choose among the
    frozen candidate list using rank-decay, with declaration order breaking ties.
    This cannot undo previous human inspection of these six development cases.
    """
    fitted = {}

    def model_for(name, training):
        key = (name, tuple(sorted(case.file_id for case in training)))
        if not training:
            raise ValueError("ACV nested validation needs at least three distinct input groups.")
        if key not in fitted:
            # Restrict the labels too, so a fitter cannot inspect held-out answers.
            fitted[key] = _fit(name, training, {case.file_id: labels[case.file_id] for case in training})
        return fitted[key]

    def prediction(name, training, heldout, universe):
        ranked, scores = _ranking(model_for(name, training), heldout)
        return {"file_id": heldout.file_id,
                "heldout_files": [case.file_id for case in universe if groups[case.file_id] == groups[heldout.file_id]],
                "training_files": [case.file_id for case in training], "faulty_car": labels[heldout.file_id],
                "ranked_cars": ranked, "score": rank_decay(ranked, labels[heldout.file_id], len(ranked)),
                "true_car_rank": ranked.index(labels[heldout.file_id]) + 1,
                "available_cars": [car for car in ranked if heldout.available[car]], "ranking_scores": scores}

    def held_out_predictions(universe):
        result = {name: [] for name in names}
        for heldout in universe:
            training = [case for case in universe if groups[case.file_id] != groups[heldout.file_id]]
            for name in names:
                result[name].append(prediction(name, training, heldout, universe))
        return result

    def summaries(predictions):
        return [{"name": name, "score": float(np.mean([row["score"] for row in predictions[name]])),
                 "ranking_metrics": ranking_metrics(predictions[name])} for name in names]

    if len(set(groups.values())) < 3:
        raise ValueError("ACV nested validation needs at least three distinct input groups.")
    predictions = held_out_predictions(cases)
    candidates = summaries(predictions)
    outer = []
    for heldout in cases:
        training = [case for case in cases if groups[case.file_id] != groups[heldout.file_id]]
        inner_predictions = held_out_predictions(training)
        inner_candidates = summaries(inner_predictions)
        selected = max(inner_candidates, key=lambda item: item["score"])
        row = prediction(selected["name"], training, heldout, cases)
        row.update({"selected_model": selected["name"], "inner_candidates": inner_candidates,
                    "inner_folds": inner_predictions[selected["name"]]})
        outer.append(row)
    nested = {"metric": "Linear rank-decay score", "score": float(np.mean([row["score"] for row in outer])),
              "ranking_metrics": ranking_metrics(outer), "folds": outer,
              "method": "Retrospective nested leave-one-input-group-out: outer case excluded from inner candidate selection and all fitting; inner rank-decay selects the candidate; declared candidate order breaks ties.",
              "limitations": ["This audits model selection, not the final model refitted on all cases.",
                              "All available cases were already inspected during development; this is not an untouched external test.",
                              f"Only {len(cases)} labelled cases in {len(set(groups.values()))} distinct selected-input groups; unknown shared acquisition may remain."]}
    return predictions, candidates, nested


def train(data_dir: Path, artifact_dir: Path) -> dict:
    data_dir, artifact_dir = Path(data_dir), Path(artifact_dir)
    labels_path = data_dir / "Train_Labels.csv"
    try:
        labels_frame = pd.read_csv(labels_path, dtype=str, keep_default_na=False)
    except (OSError, pd.errors.ParserError) as exc:
        raise ValueError(f"Cannot read ACV training labels: {exc}") from exc
    if not {"filename", "faulty_car"}.issubset(labels_frame.columns) or labels_frame.filename.duplicated().any():
        raise ValueError("ACV labels need one filename and faulty_car per unique training workbook.")
    files = sorted((data_dir / "Train").glob("*.xlsx"))
    labels = dict(zip(labels_frame.filename, labels_frame.faulty_car))
    if len(files) < 3 or {p.name for p in files} != set(labels):
        raise ValueError("ACV requires at least three labelled cases and an exact join between Train workbooks and Train_Labels.csv.")
    artifact_dir.mkdir(parents=True, exist_ok=True)
    cases, hashes = [], []
    for path in files:
        checksum = _sha(path)
        hashes.append({"file": path.name, "sha256": checksum})
        cache_path = artifact_dir / f"features-v{VERSION}-{checksum}.joblib"
        case = joblib.load(cache_path) if cache_path.exists() else load_case(path)
        if _sha(path) != checksum:
            raise ValueError(f"ACV training source {path.name} changed while its features were read.")
        if not cache_path.exists():
            joblib.dump(case, cache_path)
        case.file_id = path.name
        if labels[path.name] not in case.features.index:
            raise ValueError(f"ACV label for {path.name} does not match its header-derived car IDs.")
        if not case.available[labels[path.name]]:
            raise ValueError(f"The labelled faulty car in {path.name} has no usable cooling evidence; cannot train this feature pipeline.")
        cases.append(case)
    # Identical selected telemetry must never straddle a fold, including copies
    # saved with different workbook metadata or compression.
    group_by_file = {case.file_id: case.content_sha256 for case in cases}
    for checksum in set(group_by_file.values()):
        duplicate_labels = {labels[name] for name, value in group_by_file.items() if value == checksum}
        if len(duplicate_labels) > 1:
            raise ValueError("Identical ACV training workbooks have conflicting faulty-car labels.")
    names = _candidates()
    predictions, candidates, nested = evaluate_cases(cases, labels, group_by_file, names)
    selected = max(candidates, key=lambda item: item["score"])
    model = _fit(selected["name"], cases, labels)
    limitations = [f"Only {len(cases)} labelled case files in {len(set(group_by_file.values()))} distinct selected-input groups are available; acquisition independence is unverified."] + LIMITATIONS[1:]
    metadata = {"subsystem": "acv", "model_name": selected["name"], "trained_at": datetime.now(timezone.utc).isoformat(), "training_files": len(cases), "training_rows": sum(case.rows for case in cases), "feature_names": FEATURES, "artifact_version": VERSION,
                "training_sources": hashes + [{"file": labels_path.name, "sha256": _sha(labels_path)}],
                "validation": {"metric": "Linear rank-decay score", "score": selected["score"], "ranking_metrics": selected["ranking_metrics"], "nested_validation": nested,
                               "method": "Leave one whole selected-input group out; identical parsed workbooks excluded together; preprocessing fitted only on remaining cases. Scores select the final candidate and are not independent final-test accuracy.", "candidates": candidates, "folds": predictions[selected["name"]], "limitations": limitations},
                "limitations": limitations, "input_groups": group_by_file, "selected_model_features": COOLING_FEATURES if selected["name"] == COOLING else FEATURES,
                "preprocessing": {"validity": "Explicit Information Valid=Valid when present; exact Automatic Cooling/Full Cooling/Half Cooling modes; nonmissing control states; finite cabin and cooling-target values", "sentinels": "Observation-area channel is not used; -50 values there flagged as suspected sentinels, not diagnosed temperatures", "ranking_features": "Case-relative temperature-residual and matched-mode peer summaries; no car ID or timestamp label features", "missing_evidence": "Unavailable cars excluded from model fitting and placed at ranking tail with unknown status"}}
    token = uuid.uuid4().hex
    temporary_model = artifact_dir / f".model-{token}.tmp"
    temporary_metadata = artifact_dir / f".metadata-{token}.tmp"
    try:
        joblib.dump(model, temporary_model)
        metadata["model_sha256"] = _sha(temporary_model)
        temporary_metadata.write_text(json.dumps(metadata, indent=2, allow_nan=False), encoding="utf-8")
        temporary_model.replace(artifact_dir / "model.joblib")
        temporary_metadata.replace(artifact_dir / "metadata.json")
    finally:
        temporary_model.unlink(missing_ok=True)
        temporary_metadata.unlink(missing_ok=True)
    (artifact_dir / "validation_predictions.json").write_text(json.dumps(predictions, indent=2, allow_nan=False), encoding="utf-8")
    return metadata


def predict(path: Path, artifact_dir: Path) -> dict:
    path, artifact_dir = Path(path), Path(artifact_dir)
    model = _load_model(artifact_dir, "acv", VERSION)
    if model["feature_names"] != FEATURES or model["name"] not in (BASELINE, PAIRWISE, FOREST, BOOST, COOLING):
        raise ValueError("The frozen ACV artifact uses an incompatible ranking feature schema.")
    if model["name"] != BASELINE:
        method = "score_samples" if model["name"] == FOREST else "decision_function" if model["name"] in (PAIRWISE, COOLING) else "predict"
        if (not callable(getattr(model.get("preprocess"), "transform", None)) or
                not callable(getattr(model.get("estimator"), method, None))):
            raise ValueError("The frozen ACV artifact is missing its preprocessor or ranking estimator.")
    case = load_case(path)
    try:
        ranking, scores = _ranking(model, case)
    except (ValueError, TypeError, AttributeError, KeyError, IndexError, OverflowError) as exc:
        raise ValueError("The frozen ACV model could not rank the available measurements; check its compatible artifact and source cells.") from exc
    available_count = sum(case.available.values())
    summary = f"Car {ranking[0]} ranks first for inspection among {available_count} cars with usable cooling evidence." if available_count else "No car has usable cooling evidence; the required ordering is an uninformative identifier-order fallback."
    evidence = [{"id": "acv-coverage", "label": "Cars with usable cooling evidence", "value": f"{available_count} / {len(ranking)}", "source": path.name, "detail": "At least three valid cooling readings with numeric cabin and target values; absent evidence is not a healthy condition."}, {"id": "acv-cadence", "label": "Observed modal sampling interval", "value": case.cadence, "unit": "seconds", "source": path.name}]
    # Show observed inputs, not invented attribution: medians can tie at zero
    # while the upper tail and duration-weighted exceedance shares differ.
    for car in ranking[:9]:
        detail = case.details[car]
        evidence.append({"id": f"acv-car-{car}", "label": f"Car {car}: median cabin-minus-target residual", "value": detail["residual_median"], "unit": "raw temperature units", "source": f"{path.name} / {detail['cabin_column']} minus {detail['target_column']}", "detail": f"{detail['usable_cooling_rows']} usable cooling readings; {detail['peer_comparison_rows']} peer-comparable readings. Case-level evidence, not an onset or per-reading fault annotation."})
        for feature, label, unit, explanation in (
            ("above_target_fraction", "Cooling readings above target", "% of weighted observations", "Share of valid cooling observations with cabin measurement above the recorded cooling target, weighted by observed interval duration; long acquisition gaps are not filled."),
            ("above_peer_fraction", "Cooling residual above peers", "% of weighted observations", "Share of peer-comparable observations with cabin-minus-target residual above other cooling cars' median residual, weighted by observed interval duration. Same running modes are preferred when at least two peers exist."),
            ("peer_q90_rank", "Upper-tail peer-deviation percentile", "case percentile", f"Within-case percentile rank of the duration-weighted 90th percentile of residual above cooling peers. Raw upper-tail deviation: {detail['peer_gap_q90'] if detail['peer_gap_q90'] is not None else 'unavailable'} source temperature units; median peer deviation: {detail['peer_gap_median'] if detail['peer_gap_median'] is not None else 'unavailable'}. Equal values receive average ranks."),
        ):
            value = _finite(case.features.loc[car, feature])
            used_features = COOLING_FEATURES if model["name"] == COOLING else FEATURES
            role = "Observed model input" if feature in used_features else "Descriptive context, not a selected-model input"
            description = explanation + f" {role}: {feature}; not a probability or feature-contribution estimate."
            if value is None:
                description += " This measured feature is unavailable before model imputation; missing evidence does not mean zero risk."
            if feature == "above_target_fraction":
                warming = _finite(case.features.loc[car, "warming_fraction"])
                description += f" Stable-regime warming share (descriptive context): {warming * 100:.3f}% of weighted observations." if warming is not None else " Stable-regime warming share is unavailable."
            evidence.append({"id": f"acv-{feature}-{car}", "label": f"Car {car}: {label}", "value": value * 100 if value is not None else None, "unit": unit, "source": f"{path.name} / Car {car} cooling cabin, target and peer observations", "detail": description})
    warnings = case.warnings + LIMITATIONS[2:]
    if len(ranking) > 9:
        warnings.append("Detailed evidence is shown for the first nine ranked cars; the prediction and entity list retain every car identifier.")
    return {"subsystem": "acv", "file_id": path.name, "model_name": model["name"], "summary": summary + " Ranking is not a calibrated fault probability or verified leak diagnosis.", "prediction_rows": [{"file_id": path.name, "ranked_cars": "|".join(ranking)}], "warnings": list(dict.fromkeys(warnings)), "evidence": evidence, "series": case.series,
            "entities": [{"id": car, "label": f"Car {car}", "value": i + 1, "status": "ranked" if case.available[car] else "unavailable", "detail": f"{case.details[car]['usable_cooling_rows']} usable cooling readings; relative model score {scores[car]:.4g} (not probability)" if scores[car] is not None else "No sufficient cooling evidence. Tail position is administrative; condition unknown."} for i, car in enumerate(ranking)]}
